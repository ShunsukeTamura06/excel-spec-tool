"""統合設計書 Markdown 生成モジュール.

Workbook と ReferenceIndex から、人が読める統合設計書 (Markdown) を生成する。
LLM 注釈は別関数 (annotate_with_llm) で付与し、本モジュールは LLM 非依存とする。

docs/SPEC.ja.md §4.4 参照。
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterable

from core.diagrams import Diagram, build_diagrams
from core.external_functions import get_function, list_functions
from core.models import (
    CellFormula,
    ReferenceIndex,
    Workbook,
)

_TOP_FORMULAS_PER_SHEET = 10
_TOP_REFERENCES_GLOBAL = 20
_REFERENCE_ANALYSIS_NOTE = (
    "参照解析は静的解析です。数式参照と、VBA の静的に確定できる "
    '`Range("A1")` / `Worksheets("Sheet").Range("A1")` / '
    '`Sheets("Sheet").Cells(2, 8)` / `[Sheet!A1]` などを対象にします。'
    "加えて、OOXML から明示的に取れたグラフ系列参照とピボット元データを対象にします。"
    '`Range("A" & row)`、`Range(addr)`、`ActiveSheet`、`Selection`、'
    "`Offset` / `Resize` など実行時に決まる参照は検出対象外です。"
    "参照が 0 件でも、動的参照まで含めて影響がないとは限りません。"
)


def _md_escape(s: str) -> str:
    """Markdown テーブル中で扱いづらい文字をエスケープする (パイプと改行)."""
    return s.replace("|", "\\|").replace("\n", " ")


def _md_cell_list(values: Iterable[str], *, limit: int = 3, empty: str = "-") -> str:
    """Markdown テーブルセル内に複数値をコンパクトに並べる.

    Args:
        values: 表示する文字列列.
        limit: 表示する最大件数.
        empty: 値がない場合の表示.

    Returns:
        `<br>` 区切りの短い一覧。超過分は件数だけ示す。
    """
    items = [_md_escape(v) for v in values if v]
    if not items:
        return empty
    shown = items[:limit]
    if len(items) > limit:
        shown.append(f"... 他 {len(items) - limit} 件")
    return "<br>".join(shown)


def _section_overview(wb: Workbook) -> str:
    """## 1. 概要 を生成."""
    ext_links_str = ", ".join(wb.external_links) if wb.external_links else "なし"
    chart_count = sum(len(s.charts) for s in wb.sheets)
    pivot_count = sum(len(s.pivot_tables) for s in wb.sheets)
    lines = [
        "## 1. 概要",
        "",
        f"- ファイル名: `{wb.filename}`",
        f"- シート数: {len(wb.sheets)}",
        f"- VBAモジュール数: {len(wb.vba_modules)}",
        f"- グラフ数: {chart_count}",
        f"- ピボットテーブル数: {pivot_count}",
        f"- Power Query / 外部接続数: {len(wb.power_queries)}",
        f"- 外部リンク: {ext_links_str}",
    ]
    return "\n".join(lines)


def _section_sheet_list(wb: Workbook) -> str:
    """## 2. シート一覧 をテーブルで生成."""
    lines = [
        "## 2. シート一覧",
        "",
        "| シート名 | 行数 | 列数 | 数式数 | グラフ | ピボット | "
        "名前付き範囲 | 条件付き書式 | 用途 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    if not wb.sheets:
        lines.append("| (なし) | - | - | - | - | - | - | - | - |")
    for s in wb.sheets:
        lines.append(
            "| {name} | {rows} | {cols} | {f} | {ch} | {pt} | {nr} | {cf} | {purpose} |".format(
                name=_md_escape(s.name),
                rows=s.rows,
                cols=s.cols,
                f=len(s.formulas),
                ch=len(s.charts),
                pt=len(s.pivot_tables),
                nr=len(s.named_ranges),
                cf=len(s.conditional_formats),
                purpose=_md_escape(s.purpose) if s.purpose else "-",
            )
        )
    return "\n".join(lines)


def _pick_top_formulas(formulas: Iterable[CellFormula], n: int) -> list[CellFormula]:
    """重要そうな数式を上位 n 件選ぶ.

    現状ヒューリスティック: 参照数が多い順 → coord 昇順.
    """
    return sorted(formulas, key=lambda f: (-len(f.refs), f.coord))[:n]


def _section_sheet_details(wb: Workbook) -> str:
    """## 3. シート詳細を横断テーブル中心に生成する."""
    lines = ["## 3. シート詳細"]
    if not wb.sheets:
        lines += ["", "_(シートなし)_"]
        return "\n".join(lines)

    lines += [
        "",
        "### 3.1 業務サマリ",
        "",
        "| シート | 用途 | 利用シーン | IN | OUT | 主要計算 |",
        "|---|---|---|---|---|---|",
    ]
    for s in wb.sheets:
        lines.append(
            "| {name} | {purpose} | {scenario} | {inputs} | {outputs} | {calcs} |".format(
                name=_md_escape(s.name),
                purpose=_md_escape(s.purpose) if s.purpose else "_未設定_",
                scenario=_md_escape(s.usage_scenario) if s.usage_scenario else "-",
                inputs=_md_cell_list(s.inputs),
                outputs=_md_cell_list(s.outputs),
                calcs=_md_cell_list(s.main_calculations),
            )
        )

    lines += [
        "",
        "### 3.2 構成サマリ",
        "",
        "| シート | サイズ | 数式 | 名前付き範囲 | 入力規則 | フォーム | "
        "テーブル | 条件付き書式 | マージ |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for s in wb.sheets:
        lines.append(
            "| {name} | {shape} | {formulas} | {named} | {validations} | {controls} | "
            "{tables} | {formats} | {merged} |".format(
                name=_md_escape(s.name),
                shape=f"{s.rows}×{s.cols}",
                formulas=len(s.formulas),
                named=len(s.named_ranges),
                validations=len(s.data_validations),
                controls=len(s.form_controls),
                tables=len(s.tables),
                formats=len(s.conditional_formats),
                merged=len(s.merged_ranges),
            )
        )

    lines += [
        "",
        "### 3.3 主要数式",
    ]
    formula_rows: list[str] = []
    for s in wb.sheets:
        top = _pick_top_formulas(s.formulas, _TOP_FORMULAS_PER_SHEET)
        for f in top:
            refs_str = ", ".join(f.refs) if f.refs else "-"
            annot = f.annotation if f.annotation else "-"
            formula_rows.append(
                f"| {_md_escape(s.name)} | `{_md_escape(f.coord)}` | "
                f"`{_md_escape(f.formula)}` | {_md_escape(refs_str)} | "
                f"{_md_escape(annot)} |"
            )
    if not formula_rows:
        lines += ["", "_(数式なし)_"]
    else:
        lines += [
            "",
            f"各シート最大 {_TOP_FORMULAS_PER_SHEET} 件。",
            "",
            "| シート | セル | 数式 | 参照先 | 注釈 |",
            "|---|---|---|---|---|",
            *formula_rows,
        ]

    asset_rows: list[str] = []
    for s in wb.sheets:
        for nr in s.named_ranges:
            asset_rows.append(
                f"| {_md_escape(s.name)} | 名前付き範囲 | `{_md_escape(nr.name)}` | "
                f"`{_md_escape(nr.refers_to)}` | - |"
            )
        for cf in s.conditional_formats:
            asset_rows.append(
                f"| {_md_escape(s.name)} | 条件付き書式 | `{_md_escape(cf.range)}` | "
                f"{_md_escape(cf.rule)} | - |"
            )
        for dv in s.data_validations:
            op = f" {dv.operator}" if dv.operator else ""
            value = f"{_md_escape(dv.formula)}{_md_escape(op)}".strip() if dv.formula else "-"
            asset_rows.append(
                f"| {_md_escape(s.name)} | 入力規則 | `{_md_escape(dv.range)}` | "
                f"{_md_escape(dv.type)} | {value} |"
            )
        for fc in s.form_controls:
            target = _md_escape(fc.text or fc.name or fc.kind)
            anchor = f"`{_md_escape(fc.anchor)}`" if fc.anchor else "-"
            macro = f"`{_md_escape(fc.macro)}`" if fc.macro else "-"
            asset_rows.append(
                f"| {_md_escape(s.name)} | フォーム | {target} | {macro} | 配置: {anchor} |"
            )
        for chart in s.charts:
            target = _md_escape(chart.title or chart.name or chart.chart_type or "グラフ")
            refs = []
            for series in chart.series[:5]:
                if series.values_ref:
                    refs.append(series.values_ref)
                if series.categories_ref:
                    refs.append(series.categories_ref)
            refs_str = _md_cell_list(refs, limit=5)
            anchor = f"`{_md_escape(chart.anchor)}`" if chart.anchor else "-"
            asset_rows.append(
                f"| {_md_escape(s.name)} | グラフ | {target} | "
                f"{_md_escape(chart.chart_type) or '-'} | 系列参照: {refs_str}<br>配置: {anchor} |"
            )
        for pivot in s.pivot_tables:
            source = pivot.source_name or (
                f"{pivot.source_sheet}!{pivot.source_ref}"
                if pivot.source_sheet and pivot.source_ref
                else pivot.source_ref
            )
            source = source or "-"
            fields = (
                f"行: {_md_cell_list(pivot.row_fields, limit=4)}<br>"
                f"列: {_md_cell_list(pivot.column_fields, limit=4)}<br>"
                f"値: {_md_cell_list(pivot.value_fields, limit=4)}"
            )
            anchor = f"`{_md_escape(pivot.anchor)}`" if pivot.anchor else "-"
            asset_rows.append(
                f"| {_md_escape(s.name)} | ピボットテーブル | `{_md_escape(pivot.name)}` | "
                f"`{_md_escape(source)}` | {fields}<br>配置: {anchor} |"
            )
        for t in s.tables:
            asset_rows.append(
                f"| {_md_escape(s.name)} | Excel テーブル | `{_md_escape(t.name)}` | "
                f"`{_md_escape(t.ref)}` | ヘッダ行: {t.header_row_count} |"
            )
        if s.merged_ranges:
            shown = s.merged_ranges[:10]
            extra = len(s.merged_ranges) - len(shown)
            suffix = f" 他 {extra} 件" if extra > 0 else ""
            asset_rows.append(
                f"| {_md_escape(s.name)} | マージセル | {len(s.merged_ranges)} 件 | "
                f"{', '.join(f'`{_md_escape(m)}`' for m in shown)} |{suffix} |"
            )

    lines += ["", "### 3.4 定義・入力制約・配置物（Excel テーブル含む）"]
    if not asset_rows:
        lines += ["", "_(なし)_"]
    else:
        lines += [
            "",
            "| シート | 種別 | 対象 | 内容 | 補足 |",
            "|---|---|---|---|---|",
            *asset_rows,
        ]

    preview_blocks: list[str] = []
    for s in wb.sheets:
        block = _render_preview_block(s)
        if block:
            preview_blocks += ["", f"#### {s.name}"] + block
    if preview_blocks:
        lines += ["", "### 3.5 冒頭プレビュー", *preview_blocks]

    if wb.power_queries:
        lines += [
            "",
            "### 3.6 Power Query / 外部データ接続",
            "",
            "| 名前 | 種別 | 接続ID | 出力先 | ソース | コマンド | 信頼度 |",
            "|---|---|---|---|---|---|---|",
        ]
        for query in wb.power_queries:
            target = ""
            if query.target_sheet and query.target_name:
                target = f"{query.target_sheet} / {query.target_name}"
            elif query.target_sheet:
                target = query.target_sheet
            else:
                target = "-"
            query_name = _md_escape(query.name)
            connection_id = _md_escape(query.connection_id)
            source = _md_escape(query.source) if query.source else "-"
            command = _md_escape(query.command) if query.command else "-"
            lines.append(
                f"| {query_name} | {query.kind} | `{connection_id}` | "
                f"{_md_escape(target)} | {source} | {command} | {query.confidence} |"
            )
        lines += [
            "",
            "_Power Query の M コード本文は初期対応では解析対象外です。"
            "依存関係は、接続定義や出力先から明示的に取れる範囲だけを扱います。_",
        ]

    return "\n".join(lines)


def _render_preview_block(sheet: object) -> list[str]:
    """シート冒頭プレビューを Markdown テーブルとして描画する.

    解釈は一切加えない。生のセル値をそのまま並べるだけ。
    """
    rows = getattr(sheet, "preview_rows", []) or []
    origin = getattr(sheet, "preview_origin", "") or ""
    if not rows:
        return []

    from openpyxl.utils import get_column_letter

    n_cols = max(len(r) for r in rows)
    header = ["行"] + [get_column_letter(c + 1) for c in range(n_cols)]
    sep = ["---"] + ["---"] * n_cols

    out: list[str] = [
        "",
        f"#### 冒頭プレビュー (`{origin}`)",
        "",
        "_解釈は加えていません。シートの先頭領域に literal に並んでいる値です。_",
        "",
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(sep) + " |",
    ]
    for row_idx, row in enumerate(rows, start=1):
        cells = []
        for c in range(n_cols):
            v = row[c] if c < len(row) else None
            cells.append(_md_escape(v) if v is not None else "")
        out.append(f"| {row_idx} | " + " | ".join(cells) + " |")
    return out


def _section_vba_modules(wb: Workbook) -> str:
    """## 4. VBAモジュールを横断テーブル中心に生成する.

    プロシージャの目録のみを載せ、ソースコード本体は載せない。
    本体が必要な場合は LLM 側の `get_vba_procedure(module, name)` ツールを使う
    (設計書のサイズ削減 = Phase B-1)。
    """
    lines = ["## 4. VBAモジュール"]
    if not wb.vba_modules:
        lines += ["", "_(VBAモジュールなし)_"]
        return "\n".join(lines)

    total_lines = 0
    lines += [
        "",
        "### 4.1 モジュールサマリ",
        "",
        "| モジュール | 種別 | 行数 | プロシージャ数 |",
        "|---|---|---:|---:|",
    ]
    for m in wb.vba_modules:
        module_line_count = len(m.code.splitlines()) if m.code else 0
        total_lines += module_line_count
        lines.append(
            f"| `{_md_escape(m.name)}` | {m.type} | {module_line_count} | {len(m.procedures)} |"
        )

    proc_rows: list[str] = []
    for m in wb.vba_modules:
        for p in m.procedures:
            annot = _md_escape(p.annotation) if p.annotation else "-"
            side = _md_cell_list(p.side_effects)
            trig = _md_cell_list(p.triggers)
            calls = _md_cell_list(p.calls, limit=5)
            proc_rows.append(
                f"| `{_md_escape(m.name)}` | {p.kind} | `{_md_escape(p.name)}` | "
                f"{p.start_line}-{p.end_line} | {annot} | {side} | {trig} | {calls} |"
            )

    lines += ["", "### 4.2 プロシージャ一覧"]
    if not proc_rows:
        lines += ["", "_(プロシージャなし)_"]
    else:
        lines += [
            "",
            "| モジュール | 種別 | 名前 | 行 | 注釈 | 副作用 | 起動契機 | 呼出先 |",
            "|---|---|---|---|---|---|---|---|",
            *proc_rows,
        ]

    lines += [
        "",
        f"_合計: {len(wb.vba_modules)} モジュール / {total_lines} 行。"
        "ソースコード本体は設計書には含まれません。"
        "LLM が `get_vba_procedure(module, name)` ツールで取得します。_",
    ]
    return "\n".join(lines)


def _pick_top_references(ref_index: ReferenceIndex, n: int) -> list[tuple[str, int]]:
    """参照件数が多いキー上位 n 件を (key, count) で返す."""
    counts = [(k, len(v)) for k, v in ref_index.refs.items()]
    counts.sort(key=lambda kv: (-kv[1], kv[0]))
    return counts[:n]


def _section_references(ref_index: ReferenceIndex) -> str:
    """## 5. 参照関係（抜粋）."""
    lines = [
        "## 5. 参照関係（抜粋）",
        "",
        f"> {_REFERENCE_ANALYSIS_NOTE}",
    ]
    top = _pick_top_references(ref_index, _TOP_REFERENCES_GLOBAL)
    if not top:
        lines += ["", "_(参照なし)_"]
        return "\n".join(lines)

    lines += [
        "",
        f"参照件数の多い上位 {len(top)} 件 (参照先 → 参照元の数):",
        "",
        "| 参照先 | 件数 | 参照元 (一部) |",
        "|---|---:|---|",
    ]
    for key, count in top:
        sources = ref_index.refs.get(key, [])
        # 先頭5件だけ表示
        from_list = ", ".join(_md_escape(r.from_) for r in sources[:5])
        if len(sources) > 5:
            from_list += f", … (+{len(sources) - 5})"
        lines.append(f"| `{_md_escape(key)}` | {count} | {from_list} |")
    return "\n".join(lines)


def _mermaid_label(s: str) -> str:
    """Mermaid のノードラベル ["..."] 内で安全な文字列にする.

    ダブルクォート・バックスラッシュ・改行・パイプを除去/置換する。
    日本語などはそのまま通る。
    """
    return s.replace("\\", "\\\\").replace('"', "'").replace("\n", " ").replace("|", "/")


def _mermaid_id(prefix: str, index: int) -> str:
    """順序由来の安定 ID. ノード ID にシート名等を直接使うと記号で壊れる."""
    return f"{prefix}{index}"


_MERMAID_FENCE_OPEN = "```mermaid"
_MERMAID_FENCE_CLOSE = "```"


def _render_mermaid_graph(diagram: Diagram, direction: str = "LR") -> list[str]:
    """Diagram を mermaid `graph` 構文の行リストにして返す.

    ノード ID は順序由来 (n0, n1, ...) で安定化させ、ラベルにオリジナル名を載せる。
    エッジに weight があれば `|×N|` で併記。
    """
    lines: list[str] = [f"graph {direction}"]
    # node id 解決
    id_map: dict[str, str] = {}
    for i, n in enumerate(diagram.nodes):
        nid = _mermaid_id("n", i)
        id_map[n.id] = nid
        label = _mermaid_label(n.label or n.id)
        lines.append(f'    {nid}["{label}"]')

    for e in diagram.edges:
        src = id_map.get(e.src)
        dst = id_map.get(e.dst)
        if not src or not dst:
            continue
        if e.weight >= 2:
            lines.append(f"    {src} -->|×{e.weight}| {dst}")
        else:
            lines.append(f"    {src} --> {dst}")
    return lines


def _count_external_functions(wb: Workbook) -> Counter[str]:
    """Workbook 全体で外部関数の使用回数を集計する."""
    c: Counter[str] = Counter()
    for sheet in wb.sheets:
        for f in sheet.formulas:
            for name in f.external_functions:
                c[name] += 1
    return c


def _external_function_top_locations(wb: Workbook, name: str, limit: int = 5) -> list[str]:
    """指定外部関数の使用箇所セル座標を先頭 limit 件返す."""
    out: list[str] = []
    for sheet in wb.sheets:
        for f in sheet.formulas:
            if name in f.external_functions:
                out.append(
                    f"{sheet.name}!{f.coord.split('!', 1)[-1]}" if "!" not in f.coord else f.coord
                )
                if len(out) >= limit:
                    return out
    return out


def _section_external_functions(wb: Workbook) -> str:
    """## N. 外部関数 (Bloomberg 等) — 検出された外部 Add-In 関数の一覧と定義."""
    lines: list[str] = ["## 6. 外部関数 (Bloomberg / Refinitiv 等)"]
    counts = _count_external_functions(wb)
    if not counts:
        lines += [
            "",
            "_検出された外部 Add-In 関数はありません。_",
            "",
            "(対応ベンダー: " + ", ".join(sorted({f.vendor for f in list_functions()})) + ")",
        ]
        return "\n".join(lines)

    # サマリ
    lines += [
        "",
        f"検出された外部関数: **{len(counts)} 種類 / 使用箇所 {sum(counts.values())} 件**",
        "",
        "| 関数 | ベンダー | 使用回数 | 主要箇所 (TOP 5) | 概要 |",
        "|---|---|---:|---|---|",
    ]
    for name, cnt in counts.most_common():
        fn = get_function(name)
        vendor = fn.vendor if fn else "?"
        short = _md_escape(fn.short) if fn else "_(レジストリ未登録)_"
        locs = _external_function_top_locations(wb, name, limit=5)
        locs_str = ", ".join(f"`{_md_escape(c)}`" for c in locs) if locs else "-"
        lines.append(f"| `{name}` | {vendor} | {cnt} | {locs_str} | {short} |")

    # 各関数の詳細定義
    lines += ["", "### 関数定義"]
    for name, _ in counts.most_common():
        fn = get_function(name)
        if fn is None:
            continue
        lines += ["", f"#### `{fn.name}` ({fn.vendor})"]
        lines += ["", fn.long]
        lines += ["", f"**シグネチャ**: `{fn.signature}`"]
        if fn.params:
            lines += ["", "**引数**:", ""]
            for p in fn.params:
                req = "必須" if p.required else "任意"
                type_str = f" ({p.type})" if p.type else ""
                lines.append(f"- `{p.name}`{type_str} — {req}: {p.description}")
        if fn.returns:
            lines += ["", f"**返り値**: {fn.returns}"]
        if fn.examples:
            lines += ["", "**使用例**:", "", "```excel"]
            lines += fn.examples
            lines += ["```"]
        if fn.notes:
            lines += ["", "**注意点**:", ""]
            for note in fn.notes:
                lines.append(f"- {note}")
        if fn.doc_url:
            lines += ["", f"**公式参考**: {fn.doc_url}"]

    return "\n".join(lines)


def _section_diagrams(wb: Workbook) -> str:
    """## 7. 依存グラフ — シート依存 + VBA コールを mermaid で埋め込む.

    GitHub / VSCode / 多くの Markdown ビューアが mermaid を直接描画するため、
    ダウンロード後でも依存関係を視覚的に追える。Web UI には別途インタラクティブな
    Vue Flow タブがある。
    """
    diagrams = build_diagrams(wb)
    lines: list[str] = ["## 7. 依存グラフ"]
    lines += [
        "",
        "_GitHub / VSCode 等の mermaid 対応ビューアで描画されます._"
        " インタラクティブ表示は Web UI の「ダイアグラム」タブを参照してください。",
    ]

    # シート依存
    lines += ["", "### 7.1 シート依存"]
    sd = diagrams.sheet_deps
    if not sd.nodes:
        lines += ["", "_(シートなし)_"]
    elif not sd.edges:
        lines += ["", "_(シート間の参照なし)_"]
    else:
        lines += [
            "",
            f"ノード {len(sd.nodes)} 件 / エッジ {len(sd.edges)} 本",
            "",
            _MERMAID_FENCE_OPEN,
        ]
        lines += _render_mermaid_graph(sd, direction="LR")
        lines += [_MERMAID_FENCE_CLOSE]

    # VBA コール
    lines += ["", "### 7.2 VBA コール"]
    vc = diagrams.vba_calls
    if not vc.nodes:
        lines += ["", "_(VBA プロシージャなし)_"]
    elif not vc.edges:
        lines += ["", "_(プロシージャ間の呼び出しなし)_"]
    else:
        lines += [
            "",
            f"ノード {len(vc.nodes)} 件 / エッジ {len(vc.edges)} 本",
            "",
            _MERMAID_FENCE_OPEN,
        ]
        lines += _render_mermaid_graph(vc, direction="LR")
        lines += [_MERMAID_FENCE_CLOSE]

    return "\n".join(lines)


def _section_observations(wb: Workbook) -> str:
    """## 8. 注意点・観察事項.

    未解析リスクと LLM 注釈の付与状況を出す。
    """
    sheet_with_purpose = sum(1 for s in wb.sheets if s.purpose)
    proc_with_annot = sum(1 for m in wb.vba_modules for p in m.procedures if p.annotation)
    formulas_with_annot = sum(1 for s in wb.sheets for f in s.formulas if f.annotation)
    high_risks = [r for r in wb.analysis_risks if r.severity == "high"]
    medium_risks = [r for r in wb.analysis_risks if r.severity == "medium"]
    low_risks = [r for r in wb.analysis_risks if r.severity == "low"]

    lines = [
        "## 8. 注意点・観察事項",
        "",
        "### 8.1 未解析リスク",
        "",
        "ここに出る項目は「影響なし」と断定してはいけない箇所です。"
        "改修時は手動確認の対象にしてください。",
        "",
        f"- high: {len(high_risks)} 件",
        f"- medium: {len(medium_risks)} 件",
        f"- low: {len(low_risks)} 件",
    ]
    if not wb.analysis_risks:
        lines += ["", "_検出された未解析リスクはありません。_"]
    else:
        lines += [
            "",
            "| 重大度 | 種別 | 場所 | 根拠 | 推奨確認 |",
            "|---|---|---|---|---|",
        ]
        for risk in wb.analysis_risks[:50]:
            lines.append(
                f"| {risk.severity} | {risk.category} | `{_md_escape(risk.location)}` | "
                f"{_md_escape(risk.evidence)} | {_md_escape(risk.recommendation)} |"
            )
        if len(wb.analysis_risks) > 50:
            lines.append(
                f"| ... | ... | ... | 他 {len(wb.analysis_risks) - 50} 件 | ... |"
            )

    lines += [
        "",
        "### 8.2 LLM注釈の付与状況",
        "",
        f"- 用途が推定されたシート: {sheet_with_purpose} / {len(wb.sheets)}",
        f"- 注釈付きプロシージャ: {proc_with_annot}",
        f"- 注釈付き数式: {formulas_with_annot}",
    ]
    return "\n".join(lines)


def generate_spec(wb: Workbook, ref_index: ReferenceIndex) -> str:
    """Workbook と ReferenceIndex から統合設計書 (Markdown) を生成する.

    LLM 非依存. 注釈フィールド (SheetInfo.purpose, VbaProcedure.annotation,
    CellFormula.annotation) が埋まっていれば出力に反映される。

    Args:
        wb: 抽出済みワークブック.
        ref_index: 参照インデックス.

    Returns:
        Markdown 形式の設計書文字列 (末尾改行付き).
    """
    sections = [
        f"# 設計書: {wb.filename}",
        "",
        _section_overview(wb),
        "",
        _section_sheet_list(wb),
        "",
        _section_sheet_details(wb),
        "",
        _section_vba_modules(wb),
        "",
        _section_references(ref_index),
        "",
        _section_external_functions(wb),
        "",
        _section_diagrams(wb),
        "",
        _section_observations(wb),
        "",
    ]
    return "\n".join(sections)
